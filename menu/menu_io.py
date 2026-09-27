"""Bulk menu import/export (CSV / Excel) for tenant owners."""
import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction
from openpyxl import Workbook, load_workbook

from .models import Addon, AddonGroup, Category, MenuItem, MenuItemVariant

EXPORT_HEADERS = [
    'category_name',
    'category_name_en',
    'category_order',
    'item_name',
    'item_name_en',
    'description',
    'description_en',
    'base_price',
    'is_available',
    'item_order',
    'variants',
    'option_groups',
]

# variants: "Small=5;Large=7"
# option_groups: "Size|1|1|Small=0,Large=0;;Extras|0|3|Cheese=1,Sauce=0.5"


def _parse_bool(val) -> bool:
    if val is None or val == '':
        return True
    s = str(val).strip().lower()
    return s not in ('0', 'false', 'no', 'off', 'لا')


def _parse_variants(raw: str) -> list[tuple[str, Decimal]]:
    if not raw or not str(raw).strip():
        return []
    out = []
    for part in str(raw).split(';'):
        part = part.strip()
        if not part or '=' not in part:
            continue
        name, price = part.split('=', 1)
        try:
            out.append((name.strip(), Decimal(str(price).strip())))
        except (InvalidOperation, ValueError):
            continue
    return out


def _parse_option_groups(raw: str) -> list[dict]:
    if not raw or not str(raw).strip():
        return []
    groups = []
    for chunk in str(raw).split(';;'):
        chunk = chunk.strip()
        if not chunk:
            continue
        parts = chunk.split('|')
        if len(parts) < 4:
            continue
        gname = parts[0].strip()
        try:
            min_sel = int(parts[1].strip())
            max_sel = int(parts[2].strip())
        except ValueError:
            min_sel, max_sel = 0, 1
        options = []
        for opt in parts[3].split(','):
            opt = opt.strip()
            if not opt or '=' not in opt:
                continue
            oname, oprice = opt.split('=', 1)
            try:
                options.append({'name': oname.strip(), 'price': Decimal(oprice.strip())})
            except (InvalidOperation, ValueError):
                continue
        if gname:
            groups.append({
                'name': gname,
                'min_selection': min_sel,
                'max_selection': max(max_sel, min_sel),
                'options': options,
            })
    return groups


def build_menu_rows(tenant) -> list[list]:
    rows = [EXPORT_HEADERS]
    categories = (
        Category.objects.filter(tenant=tenant)
        .prefetch_related('items__variants', 'items__addon_groups__addons')
        .order_by('order', 'id')
    )
    for cat in categories:
        for item in cat.items.all().order_by('order', 'id'):
            variants = ';'.join(f'{v.name}={v.price}' for v in item.variants.all())
            group_parts = []
            for g in item.addon_groups.all():
                opts = ','.join(f'{a.name}={a.price}' for a in g.addons.all())
                group_parts.append(f'{g.name}|{g.min_selection}|{g.max_selection}|{opts}')
            option_groups = ';;'.join(group_parts)
            rows.append([
                cat.name,
                cat.name_en or '',
                cat.order,
                item.name,
                item.name_en or '',
                item.description or '',
                item.description_en or '',
                item.base_price,
                '1' if item.is_available else '0',
                item.order,
                variants,
                option_groups,
            ])
    return rows


def menu_export_csv(tenant) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in build_menu_rows(tenant):
        writer.writerow(row)
    return buf.getvalue().encode('utf-8-sig')


def menu_export_xlsx(tenant) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Menu'
    for row in build_menu_rows(tenant):
        ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def menu_template_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Menu'
    ws.append(EXPORT_HEADERS)
    ws.append([
        'مشروبات', 'Drinks', 0, 'عصير برتقال', 'Orange Juice',
        'طازج', 'Fresh', '3.50', '1', 0, 'Small=3;Large=5',
        'Extras|0|2|Ice=0,Mint=0.5',
    ])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _read_rows_from_upload(uploaded) -> list[dict]:
    name = (getattr(uploaded, 'name', '') or '').lower()
    raw = uploaded.read()
    rows = []

    if name.endswith('.xlsx') or name.endswith('.xlsm'):
        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        headers = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(h).strip() if h is not None else '' for h in row]
                continue
            if not headers:
                continue
            data = {headers[j]: (row[j] if j < len(row) else '') for j in range(len(headers))}
            if any(str(v).strip() for v in data.values() if v is not None):
                rows.append(data)
        return rows

    text = raw.decode('utf-8-sig', errors='replace')
    reader = csv.DictReader(io.StringIO(text))
    for row in reader:
        if any(str(v).strip() for v in row.values() if v is not None):
            rows.append(row)
    return rows


@transaction.atomic
def import_menu_from_rows(tenant, rows: list[dict]) -> dict:
    created_categories = 0
    created_items = 0
    errors = []

    cat_cache: dict[str, Category] = {}

    for idx, row in enumerate(rows, start=2):
        cat_name = (row.get('category_name') or row.get('category') or '').strip()
        item_name = (row.get('item_name') or row.get('name') or '').strip()
        if not cat_name or not item_name:
            errors.append({'row': idx, 'error': 'category_name and item_name are required'})
            continue
        try:
            base_price = Decimal(str(row.get('base_price') or '0'))
        except (InvalidOperation, ValueError):
            errors.append({'row': idx, 'error': 'invalid base_price'})
            continue

        cat_key = cat_name.lower()
        category = cat_cache.get(cat_key)
        if not category:
            try:
                cat_order = int(row.get('category_order') or 0)
            except (TypeError, ValueError):
                cat_order = 0
            category, created = Category.objects.get_or_create(
                tenant=tenant,
                name=cat_name,
                defaults={
                    'name_en': (row.get('category_name_en') or '').strip(),
                    'order': cat_order,
                },
            )
            if created:
                created_categories += 1
            cat_cache[cat_key] = category

        try:
            item_order = int(row.get('item_order') or 0)
        except (TypeError, ValueError):
            item_order = 0

        item, item_created = MenuItem.objects.get_or_create(
            tenant=tenant,
            category=category,
            name=item_name,
            defaults={
                'name_en': (row.get('item_name_en') or '').strip(),
                'description': (row.get('description') or '').strip(),
                'description_en': (row.get('description_en') or '').strip(),
                'base_price': base_price,
                'is_available': _parse_bool(row.get('is_available')),
                'order': item_order,
            },
        )
        if not item_created:
            item.name_en = (row.get('item_name_en') or item.name_en or '').strip()
            item.description = (row.get('description') or item.description or '').strip()
            item.description_en = (row.get('description_en') or item.description_en or '').strip()
            item.base_price = base_price
            item.is_available = _parse_bool(row.get('is_available'))
            item.order = item_order
            item.save()
        else:
            created_items += 1

        variants_raw = row.get('variants') or ''
        if variants_raw and item_created:
            for vname, vprice in _parse_variants(variants_raw):
                MenuItemVariant.objects.get_or_create(
                    menu_item=item, name=vname, defaults={'price': vprice},
                )

        groups_raw = row.get('option_groups') or ''
        if groups_raw and item_created:
            for g in _parse_option_groups(groups_raw):
                group = AddonGroup.objects.create(
                    menu_item=item,
                    name=g['name'],
                    min_selection=g['min_selection'],
                    max_selection=g['max_selection'],
                )
                for opt in g['options']:
                    Addon.objects.create(group=group, name=opt['name'], price=opt['price'])

    return {
        'created_categories': created_categories,
        'created_items': created_items,
        'errors': errors,
    }
